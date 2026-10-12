
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _format_worksheet(worksheet) -> None:
    """Apply consistent formatting to an Excel worksheet."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Format the header row.
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="244062",
        )
        cell.alignment = Alignment(
            vertical="center",
            horizontal="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 32

    # Set readable column widths and wrap long text.
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 14),
            45,
        )

        for cell in column_cells[1:]:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def _requirement_id_fields(item: dict) -> dict:
    """Return both IDs when available, supporting older result records."""
    old_id = item.get("old_id")
    new_id = item.get("new_id")
    primary_id = item.get("id", "")

    return {
        "Previous Requirement ID": old_id or (
            primary_id if item.get("status") == "removed" else ""
        ),
        "Revised Requirement ID": new_id or (
            primary_id if item.get("status") == "added" else ""
        ),
        "Requirement ID": primary_id,
    }


def export_impact_report(
    impact_results: list[dict],
    output_path: str | Path,
) -> Path:
    """
    Export preliminary impact analysis results to Excel.

    Includes requirement IDs, comparison details, review priorities,
    potential impact areas, and human-review indicators.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows = []

    for item in impact_results:
        row = _requirement_id_fields(item)

        row.update({
            "Status": item.get("status", ""),
            "Previous Text": item.get("old_text") or "",
            "Revised Text": item.get("new_text") or "",
            "Preliminary Priority": item.get(
                "review_priority", ""
            ),
            "Numeric Change": item.get("numeric_change", False),
            "Semantic Similarity": item.get(
                "semantic_similarity", ""
            ),
            "Human Review Required": item.get(
                "human_review_required", False
            ),
            "Potential Impact Areas": ", ".join(
                item.get("potential_impact_areas", [])
            ),
            "Reasons": " ".join(
                item.get("impact_reasons", [])
            ),
        })

        rows.append(row)

    columns = [
        "Previous Requirement ID",
        "Revised Requirement ID",
        "Requirement ID",
        "Status",
        "Previous Text",
        "Revised Text",
        "Preliminary Priority",
        "Numeric Change",
        "Semantic Similarity",
        "Human Review Required",
        "Potential Impact Areas",
        "Reasons",
    ]

    df = pd.DataFrame(rows, columns=columns)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name="Impact Analysis",
            index=False,
        )

        worksheet = writer.sheets["Impact Analysis"]
        _format_worksheet(worksheet)

    return output_path


def export_change_report(
    changes: list[dict],
    output_path: str | Path,
) -> Path:
    """
    Export requirement comparison results to Excel.

    Preserves old and revised IDs, source locations, and semantic
    matching evidence when that information is available.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows = []

    for item in changes:
        row = _requirement_id_fields(item)

        row.update({
            "Status": item.get("status", ""),
            "Previous Text": item.get("old_text") or "",
            "Revised Text": item.get("new_text") or "",
            "Previous Source": item.get("old_source") or "",
            "Previous Location": item.get("old_location") or "",
            "Revised Source": item.get("new_source") or "",
            "Revised Location": item.get("new_location") or "",
            "Semantic Similarity": item.get(
                "semantic_similarity", ""
            ),
            "Numeric Change": item.get("numeric_change", ""),
            "Human Review Required": item.get(
                "human_review_required", ""
            ),
        })

        rows.append(row)

    columns = [
        "Previous Requirement ID",
        "Revised Requirement ID",
        "Requirement ID",
        "Status",
        "Previous Text",
        "Revised Text",
        "Previous Source",
        "Previous Location",
        "Revised Source",
        "Revised Location",
        "Semantic Similarity",
        "Numeric Change",
        "Human Review Required",
    ]

    df = pd.DataFrame(rows, columns=columns)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name="Change Results",
            index=False,
        )

        worksheet = writer.sheets["Change Results"]
        _format_worksheet(worksheet)

    return output_path